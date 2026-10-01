from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(
    "/Users/Dell/Documents/Codex/2026-10-01/building-an-ai-powered-indigenous-digital/outputs"
)
DOCX = OUT / "Parakh_AI_Digital_Forensics_Presentation_Document.docx"


def set_cell_fill(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_borders(cell, color: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text: str, bold: bool = False, color: str = "000000") -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.1


def table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = t.rows[0].cells[idx]
        set_cell_fill(cell, "1F2937")
        set_cell_text(cell, header, bold=True, color="FFFFFF")
        set_cell_borders(cell)
        cell.width = Inches(widths[idx])
    for row in rows:
        cells = t.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
            set_cell_borders(cells[idx])
            if len(t.rows) % 2 == 0:
                set_cell_fill(cells[idx], "F7F9FC")
            cells[idx].width = Inches(widths[idx])
    doc.add_paragraph()


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.12


def bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    p.paragraph_format.space_after = Pt(4)


def remove_title_borders(doc: Document) -> None:
    title_style = doc.styles["Title"]
    p_pr = title_style.element.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def build() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    for style_name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        styles[style_name].font.color.rgb = RGBColor(0, 0, 0)
        styles[style_name].font.name = "Aptos Display" if style_name != "Normal" else "Aptos"
    remove_title_borders(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Parakh AI Powered Indigenous Digital Forensics Investigation Framework")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Presentation document for backend first prototype and investigator dashboard")
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(68, 68, 68)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("Prepared for project presentation and demo walkthrough").italic = True

    doc.add_paragraph()
    add_heading(doc, "Executive Summary", 1)
    add_body(
        doc,
        "Parakh is a self hosted digital forensics investigation framework designed to help "
        "investigators collect exported digital evidence, preserve integrity, analyse patterns, "
        "authenticate suspicious artifacts, correlate evidence across sources, and reconstruct an "
        "incident with source backed reasoning. The system is inspired by professional investigator "
        "suites, but it is built as an indigenous and offline capable prototype focused on Indian "
        "cybercrime patterns such as UPI fraud, Hindi and Hinglish communication, Aadhaar masking, "
        "PAN and IFSC detection, and Section 63 BSA 2023 certificate drafting.",
    )
    add_body(
        doc,
        "The current implementation has a FastAPI backend, a React investigator console, synthetic "
        "demo data, and test coverage for the core forensic guarantees. It does not perform device "
        "exploitation or lock bypass. It works only with logical exports and uploaded artifacts.",
    )

    add_heading(doc, "Problem We Are Solving", 1)
    for item in [
        "Investigators receive evidence from many places such as chats, call logs, bank statements, images, PDFs, location exports, server logs, and browser history.",
        "Manual review is slow because each source has a different format, timestamp style, and reliability level.",
        "Courts need proof that evidence was preserved correctly and that conclusions are traceable to original artifacts.",
        "AI can help, but every AI claim must cite evidence. Unsupported AI guesses are unsafe in forensic work.",
    ]:
        bullet(doc, item)

    add_heading(doc, "What We Are Building", 1)
    table(
        doc,
        ["Layer", "What it does", "Prototype status"],
        [
            ["Collection", "Uploads exported artifacts into a case and stores originals immutably by SHA 256.", "Implemented"],
            ["Integrity", "Records SHA 256 and MD5, verifies drift, and writes every action into a hash chained audit log.", "Implemented"],
            ["Parsing", "Converts source files into one normalized Event schema.", "Implemented for key demo sources"],
            ["Authentication", "Scores evidence using integrity, EXIF, PDF revision, and chat consistency signals.", "Implemented with labelled stubs for advanced detectors"],
            ["AI Analysis", "Extracts Indian entities, searches evidence, detects anomalies, and flags fraud rules.", "Implemented offline first"],
            ["Correlation", "Builds graph, timeline, phases, contradictions, and reconstruction narrative.", "Implemented"],
            ["Collaboration", "Supports comments, leads, review status, and prosecutor friendly reporting.", "Implemented baseline"],
            ["Reporting", "Generates court ready PDF with manifest, hashes, findings, timeline, audit result, and Section 63 BSA draft note.", "Implemented"],
        ],
        [1.25, 4.5, 1.55],
    )

    add_heading(doc, "Investigation Workflow", 1)
    table(
        doc,
        ["Step", "Investigator action", "System output"],
        [
            ["1", "Create a case and upload logical evidence exports.", "Case ID, evidence manifest, hashes, immutable storage path."],
            ["2", "Run parsing and extraction.", "Normalized events with timestamps, actors, content, location, and confidence."],
            ["3", "Review authenticity.", "Score from 0 to 100 and technical findings with severity and signal."],
            ["4", "Analyse entities and fraud patterns.", "Phones, UPI IDs, PAN, IFSC, masked Aadhaar, scam phrases, mule account patterns."],
            ["5", "Open graph and timeline.", "Entity links, centrality, communities, incident phases, and chronological event view."],
            ["6", "Ask the AI assistant.", "Answer with evidence ID and event ID citations. Unsupported claims are rejected."],
            ["7", "Generate report.", "Court ready PDF and Section 63 BSA draft certificate note for legal verification."],
        ],
        [0.6, 2.65, 4.05],
    )

    add_heading(doc, "Backend Features Now Added", 1)
    table(
        doc,
        ["Feature", "Endpoint", "Why it matters"],
        [
            ["Investigator dashboard", "/api/v1/cases/{case_id}/dashboard", "One API call gives the frontend evidence, risk, timeline, graph, fraud, anomaly, authenticity, audit, and collaboration data."],
            ["Cited AI assistant", "/api/v1/cases/{case_id}/assistant", "Every sentence must cite evidence and event identifiers."],
            ["Parser capabilities", "/api/v1/parsers/capabilities", "Shows supported ingestion sources and confirms logical acquisition only."],
            ["Chain of custody", "/api/v1/cases/{case_id}/chain-of-custody", "Displays audit entries and verifies the hash linked custody chain."],
            ["OSINT stub", "/api/v1/cases/{case_id}/osint", "Defines an offline safe enrichment interface without mandatory external calls."],
            ["Contradiction detection", "/api/v1/cases/{case_id}/contradictions", "Highlights gaps or conflicts such as document claims versus transaction records."],
        ],
        [1.75, 2.45, 3.1],
    )

    add_heading(doc, "Frontend Experience", 1)
    add_body(
        doc,
        "The frontend is designed as an investigator command center rather than a marketing page. "
        "It uses a dark professional visual style, evidence cards, source coverage, AI leads, graph "
        "visualization, authenticity radar, incident timeline, cited AI assistant panel, and chain "
        "of custody panel. It first tries to load the backend dashboard endpoint and falls back to "
        "synthetic demo data if the backend is not running, which makes presentations safer.",
    )
    table(
        doc,
        ["Screen area", "Purpose"],
        [
            ["Metrics", "Shows evidence count, normalized events, risk score, AI leads, and custody status."],
            ["Workflow strip", "Shows collect, preserve, parse, authenticate, correlate, and report stages."],
            ["Evidence intelligence", "Lists uploaded artifacts and source types."],
            ["Graph", "Shows relationships between actors, phone numbers, UPI IDs, and other entities."],
            ["Timeline", "Shows incident phases such as lure, credential capture, financial transfer, and movement."],
            ["AI assistant", "Explains findings only with evidence citations."],
        ],
        [2.0, 5.3],
    )

    add_heading(doc, "Demo Case", 1)
    add_body(
        doc,
        "The synthetic demo case is an investment and UPI scam. The victim receives a Hinglish "
        "WhatsApp pitch promising guaranteed returns, is pushed toward a fake KYC link and OTP "
        "flow, then sends money to a UPI handle. Bank statement rows show rapid pass through behavior "
        "from a suspected mule account. GPS and browser artifacts provide supporting context, while "
        "images and a suspicious PDF create authenticity review points.",
    )
    table(
        doc,
        ["Artifact", "Example value", "Forensic use"],
        [
            ["WhatsApp export", "Hinglish scam messages", "Intent, social engineering language, actor timeline."],
            ["UPI CSV", "arjun@upi and cashout@upi", "Money flow and mule account pattern."],
            ["Call log", "+91 phone number", "Contact evidence and entity correlation."],
            ["GPS JSON", "Location samples", "Movement and alibi support."],
            ["Browser SQLite", "Fake KYC URL", "Credential capture trail."],
            ["PDF receipt", "Multiple EOF markers", "Document manipulation suspicion."],
        ],
        [1.7, 2.0, 3.6],
    )

    add_heading(doc, "Architecture Details", 1)
    add_body(
        doc,
        "The backend is intentionally layered. API routes are thin. Services hold business logic. "
        "Parsers are plugins. Storage is content addressed. Machine learning and AI features are "
        "behind local provider interfaces so Ollama, sentence transformers, FAISS, pgvector, or "
        "cloud models can be added later without changing the API contract.",
    )
    table(
        doc,
        ["Directory", "Responsibility"],
        [
            ["app api v1", "REST endpoints and OpenAPI grouping."],
            ["app services", "Case, evidence, audit, authenticity, intelligence, analysis, and report logic."],
            ["app parsers", "Plugin registry and source specific artifact parsers."],
            ["app ml", "Entity extraction and citation enforcement helpers."],
            ["app models", "SQLAlchemy models for users, cases, evidence, events, entities, audit, leads, and reports."],
            ["app storage", "Content addressed immutable local storage."],
            ["frontend src", "React investigator console and API client."],
        ],
        [1.75, 5.55],
    )

    add_heading(doc, "Security And Legal Positioning", 1)
    for item in [
        "The system is self hosted and offline capable. No telemetry is required.",
        "Aadhaar like values are masked during extraction and display.",
        "The backend records an append only hash chained audit trail.",
        "The system labels stubs clearly and does not silently fake deepfake, OSINT, or LLM results.",
        "Generated Section 63 BSA 2023 certificate language is marked as a draft that must be verified with legal counsel.",
        "The system does not bypass locks, exploit devices, or perform invasive acquisition.",
    ]:
        bullet(doc, item)

    add_heading(doc, "How To Present This Project", 1)
    table(
        doc,
        ["Slide", "Talking point"],
        [
            ["1", "Parakh turns scattered digital evidence into a clear, court traceable investigation story."],
            ["2", "The problem is not lack of data. The problem is too much unconnected data."],
            ["3", "Our workflow is collect, preserve, parse, authenticate, analyse, correlate, reconstruct, report."],
            ["4", "The backend is the product core. The frontend is only a visual layer over documented APIs."],
            ["5", "India specific features include Hinglish text, +91 numbers, UPI, IFSC, PAN, Aadhaar masking, IST, and BSA certificate draft."],
            ["6", "AI is human in the loop. It suggests leads and summaries, but every claim must cite evidence."],
            ["7", "The demo shows an investment UPI scam from chat lure to payment movement and report generation."],
        ],
        [0.8, 6.5],
    )

    add_heading(doc, "Known Limitations And Future Work", 1)
    for item in [
        "Replace heuristic image checks with validated forensic image and deepfake detectors.",
        "Add production PostgreSQL with pgvector or FAISS for stronger multilingual semantic search.",
        "Add Celery or RQ for long running parsing and analysis jobs.",
        "Add MinIO storage adapter for large evidence stores.",
        "Add formal case membership permissions and prosecutor exhibit bundles.",
        "Add verified OSINT connectors with provider audit records and source URLs.",
        "Expand parser coverage for more phone export formats and cloud provider exports.",
    ]:
        bullet(doc, item)

    add_heading(doc, "Current Verification", 1)
    add_body(
        doc,
        "The backend was checked after the latest feature additions. The test suite passed, Ruff "
        "reported no lint issues, mypy reported no typing issues, and OpenAPI generation confirmed "
        "the new dashboard, assistant, parser capability, and chain of custody endpoints.",
    )

    doc.save(DOCX)


if __name__ == "__main__":
    build()
    print(DOCX)
